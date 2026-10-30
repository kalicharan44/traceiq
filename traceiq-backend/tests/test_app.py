from __future__ import annotations

import asyncio
import io
import json

import httpx
import pytest
import fitz
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from docx import Document as WordDocument

from app.database import get_db
from app.main import create_app
from app.models import Base, InventoryBalance, Lot, ProductionInput, ProductionOrder, Shipment, Supplier
from app.services import seed_demo_data, seed_synthetic_dataset
from app.extraction import _gemini_supply_chain_fields


def test_gemini_document_candidates_require_source_evidence(monkeypatch):
    class FakeResponse:
        def raise_for_status(self):
            return None

        def json(self):
            return {"candidates": [{"content": {"parts": [{"text": json.dumps({"fields": [
                {"field_name": "lot_identifier", "value": "RAW-778", "source_snippet": "Material lot: RAW-778", "confidence": 0.91},
                {"field_name": "supplier_name", "value": "Invented Supplier", "source_snippet": "not in the document", "confidence": 0.99},
            ]})}]}}]}

    class FakeClient:
        def __init__(self, **kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

        def post(self, *args, **kwargs):
            return FakeResponse()

    monkeypatch.setattr("app.extraction.httpx.Client", FakeClient)
    text = "Material lot: RAW-778"
    fields = _gemini_supply_chain_fields(text, [{"page_number": 2, "source_reference": None, "text": text}], "test-key")
    assert len(fields) == 1
    assert fields[0]["value"] == "RAW-778"
    assert fields[0]["page_number"] == 2
    assert fields[0]["method"] == "gemini-candidate"
    assert fields[0]["confidence"] == 0.91


def test_supply_chain_pdf_labels_extract_without_gemini(monkeypatch):
    monkeypatch.setattr("app.extraction.gemini_api_key", lambda: "")
    source = fitz.open()
    page = source.new_page()
    page.insert_text((72, 72), "Supplier ID: SUP-001\nMaterial Name: Acetaminophen API\nLot Number: RAW-1042\nProduction Order ID: PO-5001\nInput Lot ID: RAW-1042\nInput Quantity: 250\nInput Unit: kg")
    extraction = _extract_supply_chain_pdf(source.tobytes())
    values = {field["field_name"]: field["value"] for field in extraction}
    assert values["supplier_id"] == "SUP-001"
    assert values["material_name"] == "Acetaminophen API"
    assert values["lot_identifier"] == "RAW-1042"
    assert values["production_order_id"] == "PO-5001"
    assert values["input_lot_id"] == "RAW-1042"
    assert values["input_quantity"] == "250"
    assert values["input_unit"] == "kg"


def _extract_supply_chain_pdf(content):
    from app.extraction import extract_document

    return extract_document(content, "pdf").fields


class InProcessClient:
    """Run ASGI requests without opening a socket or depending on TestClient portals."""

    def __init__(self, app):
        self.transport = httpx.ASGITransport(app=app)

    def request(self, method, path, **kwargs):
        async def send():
            async with httpx.AsyncClient(transport=self.transport, base_url="http://traceiq.test") as client:
                return await asyncio.wait_for(client.request(method, path, **kwargs), timeout=8)
        return asyncio.run(send())

    def get(self, path, **kwargs):
        return self.request("GET", path, **kwargs)

    def post(self, path, **kwargs):
        return self.request("POST", path, **kwargs)

    def patch(self, path, **kwargs):
        return self.request("PATCH", path, **kwargs)


@pytest.fixture()
def client():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSession = sessionmaker(bind=engine, autoflush=False, autocommit=False, expire_on_commit=False)
    Base.metadata.create_all(bind=engine)
    with TestingSession() as db:
        seed_demo_data(db)

    app = create_app(initialize_database=False)

    def override_db():
        with TestingSession() as db:
            yield db

    app.dependency_overrides[get_db] = override_db
    test_client = InProcessClient(app)
    yield test_client, TestingSession
    app.dependency_overrides.clear()
    Base.metadata.drop_all(bind=engine)
    engine.dispose()


def test_gemini_investigation_uses_backend_key_and_validates_json(monkeypatch):
    from app import ai

    monkeypatch.setenv("GEMINI_API_KEY", "test-gemini-key")
    monkeypatch.setenv("GEMINI_MODEL", "gemini-flash-latest")
    captured: dict[str, object] = {}

    class GeminiClientStub:
        def __init__(self, **kwargs):
            captured["timeout"] = kwargs["timeout"]

        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, traceback):
            return False

        async def post(self, url, *, headers, json):
            captured.update(url=url, headers=headers, payload=json)
            body = {"candidates": [{"content": {"parts": [{"text": '{"summary":"Scope is limited to the supplied trace.","evidence":["FG-7001"],"uncertainties":[],"suggested_next_steps":["Review the linked shipment."]}'}]}}]}
            request = httpx.Request("POST", url)
            return httpx.Response(200, json=body, request=request)

    monkeypatch.setattr(ai.httpx, "AsyncClient", GeminiClientStub)
    answer = asyncio.run(ai.investigate({"finished_lots": ["FG-7001"]}, "Summarize the downstream trace."))

    assert answer.generated_by_ai is True
    assert answer.evidence == ["FG-7001"]
    assert captured["url"] == "https://generativelanguage.googleapis.com/v1beta/models/gemini-flash-latest:generateContent"
    assert captured["headers"]["X-goog-api-key"] == "test-gemini-key"
    assert captured["payload"]["generationConfig"]["responseMimeType"] == "application/json"
    assert "test-gemini-key" not in json.dumps(captured["payload"])


def test_gemini_investigation_reports_missing_key_without_network(monkeypatch):
    from app.ai import investigate

    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    answer = asyncio.run(investigate({"target_id": "RAW-1042"}, "Summarize the trace."))

    assert answer.generated_by_ai is False
    assert "GEMINI_API_KEY" in answer.summary


def test_health_and_dashboard(client):
    http, _ = client
    assert http.get("/api/v1/health").json()["status"] == "ok"
    summary = http.get("/api/v1/dashboard").json()
    assert summary["supplier_count"] == 3
    assert summary["raw_lot_count"] == 4
    assert summary["finished_lot_count"] == 4
    assert summary["open_anomaly_count"] >= 3


def test_executive_dashboard_returns_period_aware_persisted_summary(client):
    http, _ = client
    response = http.get("/api/v1/dashboard/summary", params={"period": "30d"})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["period"] == "30d"
    assert len(body["series"]) == 30
    assert body["kpis"]["suppliers"] == {"total": 3, "active": 3}
    assert body["kpis"]["material_lots"]["total"] == 8
    assert body["kpis"]["production_orders"]["completed"] == 3
    assert body["kpis"]["quality"]["critical_failures"] >= 1
    assert body["kpis"]["documents"] == {"processed": 0, "pending_review": 0, "failed": 0}
    assert body["document_processing"]["total"] == 0
    assert len(body["risks"]) == 6
    assert body["generated_at"]
    assert http.get("/api/v1/dashboard/summary", params={"period": "quarter"}).status_code == 422


def test_traceability_returns_upstream_and_downstream_scope(client):
    http, _ = client
    result = http.get("/api/v1/traceability", params={"targetId": "RAW-1042"}).json()
    assert {"FG-7001", "FG-7002"}.issubset(set(result["finished_lots"]))
    assert {"CUST-001", "CUST-002", "CUST-003"}.issubset(set(result["customer_ids"]))
    assert "SUP-001" in result["suppliers"]
    assert result["quantity_findings"]


def test_quality_check_fails_overallocated_finished_lot(client):
    http, _ = client
    result = http.post("/api/v1/quality/checks", json={"entity_type": "LOT", "entity_id": "FG-7003", "checked_by": "QA Reviewer"})
    assert result.status_code == 201
    body = result.json()
    assert body["result"] == "FAIL"
    assert any(item["code"] == "QUANTITY_OVERAGE" for item in body["findings"])
    assert http.get("/api/v1/quality/corrective-actions", params={"status": "OPEN"}).json()


def test_known_data_gaps_and_quantity_scenario_are_detected(client):
    http, _ = client
    findings = http.get("/api/v1/anomalies").json()
    types = {row["type"] for row in findings}
    assert "MISSING_INPUT_LINK" in types
    assert "QUANTITY_MISMATCH" in types
    assert "MISSING_CUSTOMER_LINK" in types

    # Test-only ground truth says shipment SHP-9005 should have reached CUST-004.
    # The operational record omits that link; tracing its source must expose the gap.
    trace = http.get("/api/v1/lots/RAW-1045/traceability").json()
    assert "CUST-004" not in trace["customer_ids"]
    assert any("SHP-9005" in item and "Customer link" in item for item in trace["missing_links"])


def test_recall_requires_approval_contains_stock_holds_shipments_and_audits(client):
    http, _ = client
    response = http.post("/api/v1/recalls", json={
        "target_lot_id": "RAW-1042",
        "reason": "Supplier reported contamination in incoming material.",
        "authorized_by": "Quality Director",
        "approval_note": "Reviewed the verified impact scope and approve activation.",
    })
    assert response.status_code == 201, response.text
    recall = response.json()
    assert recall["status"] == "ACTIVE"
    assert recall["authorized_by"] == "Quality Director"
    assert recall["affected_units"] == 1700
    assert any(action["action_type"] == "QUARANTINE_INVENTORY" for action in recall["actions"])
    assert any(action["action_type"] == "HOLD_SHIPMENT" for action in recall["actions"])

    inventory = http.get("/api/v1/inventory", params={"q": "FG-7001", "limit": 20}).json()
    assert inventory and all(row["status"] == "QUARANTINED" for row in inventory)
    shipments = http.get("/api/v1/shipments", params={"q": "SHP-9002", "limit": 20}).json()
    assert shipments[0]["status"] == "RECALL_HOLD"
    audit = http.get("/api/v1/audit-events", params={"entity_id": recall["id"]}).json()
    assert any(event["action"] == "RECALL_APPROVED_AND_ACTIVATED" for event in audit)


def test_recalled_available_stock_is_detected(client):
    http, Session = client
    response = http.post("/api/v1/recalls", json={
        "target_lot_id": "RAW-1042",
        "reason": "Confirmed supplier contamination.",
        "authorized_by": "Quality Director",
        "approval_note": "Scope reviewed and approved.",
    })
    recall_id = response.json()["id"]
    with Session() as db:
        stock = db.get(InventoryBalance, "INV-8001")
        stock.status = "AVAILABLE"  # Simulate a failed containment transition.
        db.commit()
    findings = http.get("/api/v1/anomalies").json()
    assert any(row["type"] == "RECALL_CONTAINMENT_FAILURE" and row["entity_id"] == "INV-8001" for row in findings)
    assert http.get(f"/api/v1/recalls/{recall_id}/closure-check").json()["can_close"] is False


def test_recall_closure_requires_actions_and_reconciled_quantity(client):
    http, _ = client
    response = http.post("/api/v1/recalls", json={
        "target_lot_id": "RAW-1042",
        "reason": "Supplier reported contamination in incoming material.",
        "authorized_by": "Quality Director",
        "approval_note": "Reviewed the verified impact scope and approve activation.",
    })
    recall = response.json()
    recall_id = recall["id"]
    assert http.post(f"/api/v1/recalls/{recall_id}/close", params={"actor": "QA"}).status_code == 409

    for action in recall["actions"]:
        if action["status"] != "COMPLETED":
            update = http.patch(f"/api/v1/recalls/{recall_id}/actions/{action['id']}", json={"status": "COMPLETED", "completed_by": "Recall Coordinator"})
            assert update.status_code == 200

    reconciled = http.post(f"/api/v1/recalls/{recall_id}/reconcile", json={
        "quarantined_units": 500,
        "held_or_delivered_units": 1200,
        "recorded_by": "Recall Coordinator",
    })
    assert reconciled.status_code == 200
    assert reconciled.json()["matches_scope"] is True
    closed = http.post(f"/api/v1/recalls/{recall_id}/close", params={"actor": "Quality Director"})
    assert closed.status_code == 200, closed.text
    assert closed.json()["status"] == "CLOSED"


def test_shipment_customer_gap_can_be_corrected_with_audit(client):
    http, _ = client
    corrected = http.patch("/api/v1/shipments/SHP-9005/customer", json={"customer_id": "CUST-004", "updated_by": "Data Steward"})
    assert corrected.status_code == 200
    trace = http.get("/api/v1/lots/RAW-1045/traceability").json()
    assert "CUST-004" in trace["customer_ids"]
    events = http.get("/api/v1/audit-events", params={"entity_id": "SHP-9005"}).json()
    assert any(event["action"] == "SHIPMENT_CUSTOMER_LINK_CORRECTED" for event in events)


def test_supplier_to_production_quality_gates_and_dispatch_workflow(client):
    http, _ = client
    registered = http.post("/api/v1/suppliers", json={"id": "SUP-NEW", "name": "New Qualified Source", "country": "India"})
    assert registered.status_code == 201
    assert registered.json()["status"] == "UNDER_REVIEW"

    waiting = http.post("/api/v1/quality/checks", json={"entity_type": "SUPPLIER", "entity_id": "SUP-NEW", "checked_by": "QA"})
    assert waiting.status_code == 201
    assert waiting.json()["result"] == "PENDING"

    qualified = http.post("/api/v1/quality/checks", json={"entity_type": "SUPPLIER", "entity_id": "SUP-NEW", "checked_by": "QA", "evidence_reference": "CERT-SUP-NEW-2026"})
    assert qualified.json()["result"] == "PASS"
    assert http.get("/api/v1/suppliers", params={"q": "SUP-NEW"}).json()[0]["status"] == "ACTIVE"

    intake = http.post("/api/v1/lots", json={"id": "RAW-NEW-1", "lot_type": "RAW", "name": "New API lot", "supplier_id": "SUP-NEW", "quantity": 25, "unit": "kg"})
    assert intake.status_code == 201
    raw_check = http.post("/api/v1/quality/checks", json={"entity_type": "LOT", "entity_id": "RAW-NEW-1", "checked_by": "QA", "evidence_reference": "INSP-RAW-NEW-1"})
    assert raw_check.json()["result"] == "PASS"

    production = http.post("/api/v1/production-orders", json={
        "id": "PO-NEW-1", "product_id": "PROD-001", "facility": "Pilot Facility",
        "input_lots": [{"lot_id": "RAW-NEW-1", "quantity": 5, "unit": "kg"}],
        "output_lots": [{"lot_id": "FG-NEW-1", "name": "Pilot finished lot", "quantity": 100, "unit": "units"}],
    })
    assert production.status_code == 201, production.text
    assert http.get("/api/v1/lots/FG-NEW-1").json()["status"] == "PENDING_QUALITY"
    fg_check = http.post("/api/v1/quality/checks", json={"entity_type": "LOT", "entity_id": "FG-NEW-1", "checked_by": "QA"})
    assert fg_check.json()["result"] == "PASS"

    receipt = http.post("/api/v1/inventory/receipts", json={"id": "INV-NEW-1", "lot_id": "FG-NEW-1", "warehouse_id": "WH-01", "warehouse_name": "Pilot Warehouse", "quantity": 100, "unit": "units"})
    assert receipt.status_code == 201
    dispatch = http.post("/api/v1/shipments", json={"id": "SHP-NEW-1", "lot_id": "FG-NEW-1", "customer_id": "CUST-001", "quantity": 40, "unit": "units", "destination": "Chicago, IL"})
    assert dispatch.status_code == 201
    assert dispatch.json()["status"] == "IN_TRANSIT"


def test_failed_production_gate_is_atomic(client):
    http, _ = client
    result = http.post("/api/v1/production-orders", json={
        "id": "PO-INVALID", "product_id": "PROD-001",
        "input_lots": [{"lot_id": "RAW-NOT-FOUND", "quantity": 1, "unit": "kg"}],
        "output_lots": [{"lot_id": "FG-INVALID", "name": "Should not exist", "quantity": 10}],
    })
    assert result.status_code == 409
    orders = http.get("/api/v1/production-orders", params={"limit": 200}).json()
    assert all(order["id"] != "PO-INVALID" for order in orders)


def test_failed_recall_simulation_does_not_partially_create_case(client):
    http, _ = client
    failed = http.post("/api/v1/recalls", json={
        "target_lot_id": "RAW-NOT-FOUND",
        "reason": "A test request for a missing lot.",
        "authorized_by": "Quality Director",
        "approval_note": "This should fail before any recall state changes.",
    })
    assert failed.status_code == 404
    assert http.get("/api/v1/recalls").json() == []
    assert all(row["status"] == "AVAILABLE" for row in http.get("/api/v1/inventory").json())


def test_non_lot_anomaly_record_resolves_to_related_lot(client):
    http, _ = client
    response = http.get("/api/v1/traceability", params={"targetId": "PO-5003"})
    assert response.status_code == 200, response.text
    assert response.json()["target_id"] == "FG-7003"
    assert any("RAW-GHOST-9001" in item for item in response.json()["missing_links"])


def test_synthetic_seed_has_requested_scale_and_planted_evidence_gaps():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    TestingSession = sessionmaker(bind=engine, autoflush=False, autocommit=False, expire_on_commit=False)
    Base.metadata.create_all(bind=engine)
    try:
        with TestingSession() as db:
            counts = seed_synthetic_dataset(db, seed=20261022)
            assert counts == {"suppliers": 500, "raw_lots": 10000, "production_orders": 5000, "finished_lots": 8000, "shipments": 20000}
            assert db.scalar(select(func.count()).select_from(Supplier)) == 500
            assert db.scalar(select(func.count()).select_from(Lot)) == 18000
            assert db.scalar(select(func.count()).select_from(ProductionOrder)) == 5000
            assert db.scalar(select(func.count()).select_from(Shipment)) == 20000
            assert db.scalar(select(ProductionInput).where(ProductionInput.production_order_id == "PO-00001", ProductionInput.input_lot_id == "RAW-GHOST-0001"))
            assert db.get(Lot, "FG-00003").quantity == 10
            assert db.get(Shipment, "SHP-00001").quantity == 50
            assert db.get(Shipment, "SHP-00002").customer_id is None
    finally:
        Base.metadata.drop_all(bind=engine)
        engine.dispose()


def test_private_document_upload_extract_review_link_and_validation(client, tmp_path, monkeypatch):
    # Keep this workflow test deterministic and independent from a configured
    # external Gemini service; candidate extraction has its own adapter path.
    monkeypatch.setenv("GEMINI_API_KEY", "")
    http, TestingSession = client
    monkeypatch.setenv("PRIVATE_STORAGE_DIR", str(tmp_path / "private-documents"))
    monkeypatch.setattr("app.document_controller.SessionLocal", TestingSession)

    source = WordDocument()
    source.add_paragraph("ACME Corporation")
    source.add_paragraph("Invoice Number: INV-2026-071")
    source.add_paragraph("Invoice Date: 10/29/2026")
    source.add_paragraph("Purchase Order: PO-2026-000001")
    source.add_paragraph("Total: $128.50")
    content = io.BytesIO()
    source.save(content)
    original = content.getvalue()
    multipart = {"file": ("invoice.docx", original, "application/octet-stream")}

    uploaded = http.post("/api/v1/documents", data={"category": "SUPPLIER_INVOICE", "uploaded_by": "QA Reviewer"}, files=multipart)
    assert uploaded.status_code == 201, uploaded.text
    document = uploaded.json()
    assert document["status"] == "UPLOADED"
    assert "storage_key" not in document
    listed = http.get("/api/v1/documents")
    assert listed.status_code == 200
    assert listed.json()["items"][0]["id"] == document["id"]

    duplicate = http.post("/api/v1/documents", data={"category": "SUPPLIER_INVOICE"}, files=multipart)
    assert duplicate.status_code == 409
    preview = http.get(f"/api/v1/documents/{document['id']}/content")
    assert preview.content == original
    assert "private-documents" not in preview.headers.get("content-disposition", "")

    queued = http.post(f"/api/v1/documents/{document['id']}/process", data={"actor": "QA Reviewer"})
    assert queued.status_code == 202
    status = http.get(f"/api/v1/documents/{document['id']}/processing-status").json()
    assert status["status"] == "REVIEW_REQUIRED"
    extraction = http.get(f"/api/v1/documents/{document['id']}/extraction").json()
    assert {field["field_name"] for field in extraction["fields"]} >= {"invoice_number", "invoice_date", "total_amount"}
    assert all(field["confidence"] is None for field in extraction["fields"] if field["method"] == "deterministic-pattern")

    reviewed = http.patch(f"/api/v1/documents/{document['id']}/review", json={
        "decision": "VERIFIED", "reviewer": "QA Reviewer", "note": "Compared extracted values against the source file.",
        "reviewed_fields": {"supplier_name": "ACME Corporation"},
    })
    assert reviewed.status_code == 200, reviewed.text
    linked = http.post(f"/api/v1/documents/{document['id']}/links", json={"entity_type": "SUPPLIER", "entity_id": "SUP-001", "created_by": "QA Reviewer"})
    assert linked.status_code == 201, linked.text
    checks = http.post(f"/api/v1/documents/{document['id']}/run-checks", data={"run_by": "QA Reviewer"})
    assert checks.status_code == 201, checks.text
    assert checks.json()["status"] == "PASS"
    assert http.get(f"/api/v1/checks/runs/{checks.json()['id']}").json()["findings"]

