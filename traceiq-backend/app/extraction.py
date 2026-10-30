"""Safe file validation and deterministic text extraction for TraceIQ documents."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from typing import Any

import fitz
from docx import Document as WordDocument
from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError
import httpx
from pydantic import BaseModel, Field, ValidationError
from app.config import gemini_api_key, gemini_model


EXTRACTOR_VERSION = "traceiq-extractor-v1"
ALLOWED_SUFFIXES = {".pdf", ".png", ".jpg", ".jpeg", ".docx", ".csv", ".xlsx"}
TYPE_MEDIA = {
    "pdf": "application/pdf",
    "png": "image/png",
    "jpeg": "image/jpeg",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
}


@dataclass
class ExtractionResult:
    text: str
    pages: list[dict[str, Any]]
    fields: list[dict[str, Any]]
    warnings: list[str]
    detected_category: str
    extractor_version: str = EXTRACTOR_VERSION


def _extension(filename: str) -> str:
    clean = filename.replace("\\", "/").split("/")[-1]
    return "." + clean.rsplit(".", 1)[-1].lower() if "." in clean else ""


def _safe_zip(data: bytes) -> list[str]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 5000:
                raise ValueError("Office document contains too many archive entries.")
            total_size = sum(info.file_size for info in infos)
            if total_size > 100 * 1024 * 1024:
                raise ValueError("Office document expands beyond the 100 MB safety limit.")
            if any(info.file_size > 1_000_000 and info.file_size / max(info.compress_size, 1) > 200 for info in infos):
                raise ValueError("Office document has an unsafe compression ratio.")
            names = [info.filename for info in infos]
            lowered = {name.lower() for name in names}
            if any("vbaproject.bin" in name for name in lowered):
                raise ValueError("Macro-enabled Office documents are not accepted.")
            return names
    except zipfile.BadZipFile as exc:
        raise ValueError("The uploaded Office document is not a valid ZIP-based file.") from exc


def detect_file_type(filename: str, data: bytes, max_bytes: int, max_pages: int) -> str:
    if not data:
        raise ValueError("The uploaded file is empty.")
    if len(data) > max_bytes:
        raise ValueError(f"File exceeds the configured {max_bytes // (1024 * 1024)} MB upload limit.")
    suffix = _extension(filename)
    if suffix not in ALLOWED_SUFFIXES:
        raise ValueError("Unsupported file format. Upload PDF, PNG, JPG, JPEG, DOCX, CSV, or XLSX.")

    if data.startswith(b"%PDF-"):
        file_type = "pdf"
        try:
            with fitz.open(stream=data, filetype="pdf") as pdf:
                if pdf.is_encrypted:
                    raise ValueError("Password-protected PDFs cannot be processed.")
                if len(pdf) < 1:
                    raise ValueError("The PDF contains no pages.")
                if len(pdf) > max_pages:
                    raise ValueError(f"PDF exceeds the configured {max_pages}-page limit.")
        except fitz.FileDataError as exc:
            raise ValueError("The uploaded PDF is damaged or invalid.") from exc
    elif data.startswith(b"\x89PNG\r\n\x1a\n"):
        file_type = "png"
    elif data.startswith(b"\xff\xd8\xff"):
        file_type = "jpeg"
    elif zipfile.is_zipfile(io.BytesIO(data)):
        names = {name.lower() for name in _safe_zip(data)}
        if "word/document.xml" in names and suffix == ".docx":
            file_type = "docx"
        elif "xl/workbook.xml" in names and suffix == ".xlsx":
            file_type = "xlsx"
        else:
            raise ValueError("The file contents do not match a supported DOCX or XLSX format.")
    elif suffix == ".csv":
        try:
            sample = data[:16_384].decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV files must use UTF-8 text encoding.") from exc
        if "\x00" in sample:
            raise ValueError("The uploaded CSV is not plain UTF-8 text.")
        try:
            csv.Sniffer().sniff(sample[:4096], delimiters=",;\t|")
        except csv.Error as exc:
            raise ValueError("The CSV file does not contain a recognizable delimiter.") from exc
        file_type = "csv"
    else:
        raise ValueError("The uploaded file contents do not match a supported document type.")

    expected_suffixes = {"pdf": {".pdf"}, "png": {".png"}, "jpeg": {".jpg", ".jpeg"}, "docx": {".docx"}, "xlsx": {".xlsx"}, "csv": {".csv"}}
    if suffix not in expected_suffixes[file_type]:
        raise ValueError("The file extension does not match its detected contents.")
    if file_type in {"png", "jpeg"}:
        try:
            with Image.open(io.BytesIO(data)) as image:
                image.verify()
        except (UnidentifiedImageError, OSError) as exc:
            raise ValueError("The uploaded image is damaged or invalid.") from exc
    return file_type


def _ocr_image(image: Image.Image) -> tuple[str, list[dict[str, Any]], str | None]:
    try:
        import pytesseract
        from pytesseract import Output

        data = pytesseract.image_to_data(image, output_type=Output.DICT, timeout=30)
    except ImportError:
        return "", [], "OCR support is not installed. Install pytesseract and the Tesseract executable to process scanned pages."
    except Exception as exc:  # Tesseract missing, timeout, or corrupt image should not verify the document.
        name = type(exc).__name__
        if "TesseractNotFound" in name:
            return "", [], "Tesseract is not installed or is not on PATH; this scanned page needs manual review."
        if "Timeout" in name:
            return "", [], "OCR exceeded the per-page time limit; retry processing or review manually."
        return "", [], f"OCR could not process this page ({name}); review it manually."

    words: list[dict[str, Any]] = []
    lines: dict[tuple[str, str, str], list[str]] = {}
    for index, raw_word in enumerate(data.get("text", [])):
        word = str(raw_word).strip()
        if not word:
            continue
        box = {
            "text": word,
            "bbox": [int(data["left"][index]), int(data["top"][index]), int(data["width"][index]), int(data["height"][index])],
        }
        words.append(box)
        key = (str(data["block_num"][index]), str(data["par_num"][index]), str(data["line_num"][index]))
        lines.setdefault(key, []).append(word)
    return "\n".join(" ".join(parts) for parts in lines.values()), words, None


def _page_record(number: int | None, text: str, method: str, words: list[dict[str, Any]] | None = None, reference: str | None = None) -> dict[str, Any]:
    return {"page_number": number, "source_reference": reference, "text": text, "method": method, "words": words or []}


def _extract_pdf(data: bytes, warnings: list[str]) -> list[dict[str, Any]]:
    pages: list[dict[str, Any]] = []
    with fitz.open(stream=data, filetype="pdf") as pdf:
        for index, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()
            if len(re.sub(r"\s", "", text)) >= 30:
                pages.append(_page_record(index, text, "pdf-text"))
                continue
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.open(io.BytesIO(pixmap.tobytes("png")))
            ocr_text, words, warning = _ocr_image(image)
            if warning:
                warnings.append(f"Page {index}: {warning}")
            pages.append(_page_record(index, ocr_text or text, "tesseract-ocr" if ocr_text else "pdf-text", words))
    return pages


def _extract_docx(data: bytes) -> list[dict[str, Any]]:
    doc = WordDocument(io.BytesIO(data))
    pieces: list[str] = []
    for index, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            pieces.append(f"[Paragraph {index}] {text}")
    for table_index, table in enumerate(doc.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            pieces.append(f"[Table {table_index}, row {row_index}] " + " | ".join(cells))
    return [_page_record(None, "\n".join(pieces), "docx-text", reference="DOCX paragraphs and tables")]


def _extract_csv(data: bytes) -> list[dict[str, Any]]:
    text = data.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    lines = []
    for row_index, row in enumerate(reader, start=1):
        if row_index > 5000:
            raise ValueError("CSV exceeds the 5,000-row extraction limit.")
        if any(str(cell).strip() for cell in row):
            lines.append(f"[Row {row_index}] " + " | ".join(str(cell).strip() for cell in row))
    return [_page_record(None, "\n".join(lines), "csv-parser", reference="CSV rows")]


def _extract_xlsx(data: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pages: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            lines = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if row_index > 5000:
                    raise ValueError("Each spreadsheet sheet is limited to 5,000 rows.")
                values = [str(value).strip() if value is not None else "" for value in row]
                if any(values):
                    lines.append(f"[Row {row_index}] " + " | ".join(values))
            pages.append(_page_record(None, "\n".join(lines), "xlsx-parser", reference=f"Sheet: {sheet.title}"))
    finally:
        workbook.close()
    return pages


def _classify(text: str) -> str:
    sample = text[:40_000].lower()
    if any(term in sample for term in ("certificate of analysis", "coa number", "laboratory results")):
        return "CERTIFICATE_OF_ANALYSIS"
    if any(term in sample for term in ("goods receipt note", "goods received note", "grn number")):
        return "GOODS_RECEIPT_NOTE"
    if any(term in sample for term in ("proof of delivery", "delivery confirmation", "recipient acknowledgement")):
        return "PROOF_OF_DELIVERY"
    if "invoice" in sample and any(term in sample for term in ("invoice no", "invoice number", "tax amount", "subtotal")):
        return "SUPPLIER_INVOICE"
    if "purchase order" in sample or re.search(r"\bpo\s*(number|no\.?|#)", sample):
        return "PURCHASE_ORDER"
    if any(term in sample for term in ("batch manufacturing record", "material consumption report")):
        return "BATCH_MANUFACTURING_RECORD"
    if any(term in sample for term in ("inspection report", "quality inspection")):
        return "MATERIAL_INSPECTION_REPORT"
    return "OTHER"


FIELD_PATTERNS = [
    ("supplier_id", re.compile(r"^supplier\s*(?:id|code)\s*[:#-]?\s*([A-Z0-9][A-Z0-9/-]{1,})", re.I)),
    ("supplier_name", re.compile(r"^supplier\s*name\s*[:#-]?\s*(.{2,100})", re.I)),
    ("material_name", re.compile(r"^(?:material|product)\s*(?:name)?\s*[:#-]\s*(.{2,100})", re.I)),
    ("production_order_id", re.compile(r"^production\s*order\s*(?:id|number|no\.?|#)?\s*[:#-]?\s*(PO[-/ ]?[A-Z0-9-]+)", re.I)),
    ("input_lot_id", re.compile(r"^input\s*lot\s*(?:id|number|no\.?|#)?\s*[:#-]?\s*([A-Z0-9][A-Z0-9/-]{2,})", re.I)),
    ("output_lot_id", re.compile(r"^output\s*lot\s*(?:id|number|no\.?|#)?\s*[:#-]?\s*([A-Z0-9][A-Z0-9/-]{2,})", re.I)),
    ("input_quantity", re.compile(r"^input\s*quantity\s*[:#-]\s*([0-9][0-9,]*(?:\.\d+)?)", re.I)),
    ("output_quantity", re.compile(r"^output\s*quantity\s*[:#-]\s*([0-9][0-9,]*(?:\.\d+)?)", re.I)),
    ("input_unit", re.compile(r"^input\s*unit\s*[:#-]\s*(kg|g|lb|lbs|liters?|l|units?|cases?|pallets?)\b", re.I)),
    ("output_unit", re.compile(r"^output\s*unit\s*[:#-]\s*(kg|g|lb|lbs|liters?|l|units?|cases?|pallets?)\b", re.I)),
    ("quantity", re.compile(r"^(?:quantity|qty)(?:\s+(?:received|ordered|consumed|shipped|produced))?\s*[:#-]?\s*([0-9][0-9,]*(?:\.\d+)?)", re.I)),
    ("unit", re.compile(r"^unit\s*[:#-]?\s*(kg|g|lb|lbs|liters?|l|units?|cases?|pallets?)\b", re.I)),
    ("receipt_reference", re.compile(r"^(?:goods\s*receipt|receipt|GRN)\s*(?:reference|number|no\.?|#|id)?\s*[:#-]?\s*(GRN[-/ ]?[A-Z0-9-]+)", re.I)),
    ("shipment_reference", re.compile(r"^shipment\s*(?:reference|number|no\.?|#|id)?\s*[:#-]?\s*(SHP[-/ ]?[A-Z0-9-]+)", re.I)),
    ("inventory_reference", re.compile(r"^inventory\s*(?:reference|number|no\.?|#|id)?\s*[:#-]?\s*(INV[-/ ]?[A-Z0-9-]+)", re.I)),
    ("customer_reference", re.compile(r"^customer\s*(?:reference|number|no\.?|#|id)?\s*[:#-]?\s*(CUST[-/ ]?[A-Z0-9-]+)", re.I)),
    ("invoice_number", re.compile(r"^(?:invoice\s*(?:number|no\.?|#))\s*[:#-]?\s*([A-Z0-9][A-Z0-9/-]{2,})", re.I)),
    ("document_number", re.compile(r"^(?:document\s*(?:number|no\.?|#)|certificate\s*(?:number|no\.?|#)|grn\s*(?:number|no\.?|#))\s*[:#-]?\s*([A-Z0-9][A-Z0-9/-]{2,})", re.I)),
    ("purchase_order_reference", re.compile(r"^(?:purchase\s*order|\bPO)\s*(?:reference|number|no\.?|#)?\s*[:#-]?\s*(PO[-/ ]?[A-Z0-9-]+)", re.I)),
    ("lot_identifier", re.compile(r"^(?:lot|batch)\s*(?:identifier|number|no\.?|#)?\s*[:#-]\s*([A-Z0-9][A-Z0-9/-]{2,})", re.I)),
    ("invoice_date", re.compile(r"^(?:invoice\s*)?date\s*[:#-]?\s*(\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}|[A-Za-z]{3,9}\s+\d{1,2},?\s+\d{4})", re.I)),
    ("total_amount", re.compile(r"^(?:invoice\s*)?total\s*[:#-]?\s*([$€£₹]?\s?\d[\d,]*(?:\.\d{2})?)", re.I)),
]

TABLE_FIELD_LABELS = [
    ("supplier_id", re.compile(r"^supplier\s*(?:id|code)$", re.I)),
    ("supplier_name", re.compile(r"^supplier\s*name$", re.I)),
    ("material_name", re.compile(r"^material\s*name$", re.I)),
    ("product_name", re.compile(r"^product\s*name$", re.I)),
    ("production_order_id", re.compile(r"^production\s*order\s*(?:id|number)$", re.I)),
    ("input_lot_id", re.compile(r"^input\s*lot\s*id$", re.I)),
    ("output_lot_id", re.compile(r"^output\s*lot\s*id$", re.I)),
    ("finished_lot_id", re.compile(r"^finished\s*lot\s*id$", re.I)),
    ("lot_identifier", re.compile(r"^(?:lot\s*(?:number|id)|batch\s*(?:number|id))$", re.I)),
    ("purchase_order_reference", re.compile(r"^purchase\s*order$", re.I)),
    ("receipt_reference", re.compile(r"^receipt\s*reference$", re.I)),
    ("shipment_reference", re.compile(r"^shipment\s*id$", re.I)),
    ("inventory_reference", re.compile(r"^inventory\s*reference$", re.I)),
    ("customer_reference", re.compile(r"^customer\s*id$", re.I)),
    ("customer_name", re.compile(r"^customer\s*name$", re.I)),
    ("quantity", re.compile(r"^quantity\s+(?:received|ordered|consumed|shipped|produced)$", re.I)),
    ("input_quantity", re.compile(r"^input\s*quantity$", re.I)),
    ("output_quantity", re.compile(r"^output\s*quantity$", re.I)),
    ("unit", re.compile(r"^unit$", re.I)),
    ("input_unit", re.compile(r"^input\s*unit$", re.I)),
    ("output_unit", re.compile(r"^output\s*unit$", re.I)),
    ("warehouse", re.compile(r"^warehouse$", re.I)),
    ("destination", re.compile(r"^destination$", re.I)),
    ("shipment_status", re.compile(r"^shipment\s*status$", re.I)),
    ("facility", re.compile(r"^facility$", re.I)),
]


def _rule_fields(pages: list[dict[str, Any]]) -> list[dict[str, Any]]:
    fields: list[dict[str, Any]] = []
    for field_name, pattern in FIELD_PATTERNS:
        for page in pages:
            for line in page["text"].splitlines():
                match = pattern.search(line)
                if not match:
                    continue
                value = match.group(1).strip()
                fields.append({
                    "field_name": field_name,
                    "value": value,
                    "page_number": page["page_number"],
                    "source_reference": page.get("source_reference") or (f"Page {page['page_number']}" if page["page_number"] else None),
                    "source_snippet": line.strip()[:500],
                    "bbox": None,
                    "method": "deterministic-pattern",
                    "confidence": None,
                    "review_status": "UNREVIEWED",
                })
                break
            else:
                continue
            break
    for page in pages:
        lines = [line.strip() for line in page["text"].splitlines()]
        for index, label in enumerate(lines[:-1]):
            value = lines[index + 1].strip()
            if not label or not value or len(value) > 180:
                continue
            for field_name, pattern in TABLE_FIELD_LABELS:
                if not pattern.fullmatch(label):
                    continue
                if field_name.endswith("quantity"):
                    amount = re.match(r"^([0-9][0-9,]*(?:\.\d+)?)", value)
                    if not amount:
                        continue
                    value = amount.group(1).replace(",", "")
                if any(field["field_name"] == field_name and str(field["value"]) == value for field in fields):
                    break
                fields.append({
                    "field_name": field_name,
                    "value": value,
                    "page_number": page["page_number"],
                    "source_reference": page.get("source_reference") or (f"Page {page['page_number']}" if page["page_number"] else None),
                    "source_snippet": f"{label}\n{value}",
                    "bbox": None,
                    "method": "deterministic-label-value",
                    "confidence": None,
                    "review_status": "UNREVIEWED",
                })
                break
    return fields


def extract_document(data: bytes, file_type: str) -> ExtractionResult:
    warnings: list[str] = []
    if file_type == "pdf":
        pages = _extract_pdf(data, warnings)
    elif file_type == "docx":
        pages = _extract_docx(data)
    elif file_type == "csv":
        pages = _extract_csv(data)
    elif file_type == "xlsx":
        pages = _extract_xlsx(data)
    elif file_type in {"png", "jpeg"}:
        with Image.open(io.BytesIO(data)) as image:
            text, words, warning = _ocr_image(image.convert("RGB"))
        if warning:
            warnings.append(warning)
        pages = [_page_record(1, text, "tesseract-ocr" if text else "image", words)]
    else:
        raise ValueError("No extractor is registered for this file type.")
    full_text = "\n\n".join(page["text"] for page in pages if page["text"]).strip()
    if not full_text:
        warnings.append("No readable text was extracted. Review the source document manually.")
    fields = _rule_fields(pages)
    # LLM output is a candidate only. Validate the JSON shape and attach evidence
    # from the locally extracted text; never use model supplied citations as fact.
    api_key = gemini_api_key()
    if api_key and full_text:
        try:
            fields.extend(_gemini_supply_chain_fields(full_text, pages, api_key))
        except Exception as exc:
            warnings.append(f"AI field suggestions unavailable ({type(exc).__name__}); deterministic extraction is shown.")
    elif not api_key:
        warnings.append("GEMINI_API_KEY is not configured; only deterministic document fields are available.")
    return ExtractionResult(full_text, pages, fields, warnings, _classify(full_text))


class _Candidate(BaseModel):
    field_name: str = Field(min_length=2, max_length=80)
    value: str | float | None
    source_snippet: str = Field(min_length=1, max_length=500)
    confidence: float = Field(ge=0, le=1)


class _CandidateList(BaseModel):
    fields: list[_Candidate] = Field(max_length=100)


def _gemini_supply_chain_fields(text: str, pages: list[dict[str, Any]], api_key: str) -> list[dict[str, Any]]:
    model = gemini_model()
    if not re.fullmatch(r"[A-Za-z0-9._-]+", model):
        raise ValueError("Invalid configured Gemini model name")
    prompt = ("Extract only explicitly stated supply-chain values: material/finished lot IDs and names, supplier IDs/names, "
              "production order IDs, input/output lot relationships, quantities/units, shipment, inventory, receipt and customer references. "
              "Do not infer missing values. Return JSON {fields:[{field_name,value,source_snippet,confidence}]}. "
              "Every source_snippet must be copied verbatim from the supplied text. Candidate values are untrusted.\nDOCUMENT TEXT:\n" + text[:120_000])
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
    with httpx.Client(timeout=httpx.Timeout(35.0, connect=5.0)) as client:
        response = client.post(url, headers={"X-goog-api-key": api_key}, json={"contents":[{"parts":[{"text":prompt}]}],"generationConfig":{"temperature":0,"maxOutputTokens":2500,"responseMimeType":"application/json"}})
        response.raise_for_status()
    payload = response.json()
    raw = payload["candidates"][0]["content"]["parts"][0]["text"]
    parsed = _CandidateList.model_validate_json(raw)
    output = []
    for item in parsed.fields:
        snippet = item.source_snippet.strip()
        if snippet not in text:
            continue
        page = next((p for p in pages if snippet in p["text"]), None)
        output.append({"field_name":item.field_name,"value":item.value,"page_number":page["page_number"] if page else None,
                       "source_reference":(page.get("source_reference") or f"Page {page['page_number']}" if page and page.get("page_number") else (page.get("source_reference") if page else None)),
                       "source_snippet":snippet,"bbox":None,"method":"gemini-candidate","confidence":item.confidence,"review_status":"UNREVIEWED"})
    return output
